from docx import Document

def fill_docx(template_path: str, output_path: str, values: dict):
    doc = Document(template_path)
    print(f"🔧 Заполняю поля: {values}")

    def replace_in_paragraph(p):
        full_text = p.text
        changed = False
        for key, val in values.items():
            placeholder = "{{" + key + "}}"
            if placeholder in full_text:
                full_text = full_text.replace(placeholder, str(val))
                changed = True
        if changed:
            print(f"✅ Заменено в параграфе: {p.text} -> {full_text}")
            p.runs[0].text = full_text
            for run in p.runs[1:]:
                run.text = ""
        elif "{{" in p.text:
            print(f"⚠️ Не нашёл совпадений для: {p.text}")

    for p in doc.paragraphs:
        replace_in_paragraph(p)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    replace_in_paragraph(p)

    doc.save(output_path)
    print(f"💾 Сохранено: {output_path}")
